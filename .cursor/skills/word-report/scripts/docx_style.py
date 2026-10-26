"""Shared Word styling for user-requested reports (PDT-3483 exemplar)."""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as Rel
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLUE = "0969DA"
TEXT = "1F2328"
MUTED = "656D76"
GREEN = "1A7F37"
HEADER_FILL = "F6F8FA"
ROW_ALT = "F6F8FA"
BORDER = "D0D7DE"
PASSED_FILL = "DAFBE1"


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    return doc


def set_run_font(run, size=11, bold=False, color=TEXT, italic=False):
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(p, before=0, after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_hyperlink(paragraph, text, url, size=11, bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, Rel.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rPr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz_cs)
    if bold:
        rPr.append(OxmlElement("w:b"))
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_mixed(paragraph, text, size=11, color=TEXT):
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=(i % 2 == 1), color=color)


def add_heading(doc, text, level):
    sizes = {1: 26, 2: 18, 3: 14}
    before = {1: 0, 2: 18, 3: 16}
    after = {1: 6, 2: 10, 3: 8}
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before[level], after=after[level], line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=sizes[level], bold=True, color=TEXT)
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=10, line=1.0)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BORDER)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    return p


def shade_cell(cell, fill):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tc_pr.append(borders)


def set_cell_margins(cell, dxa=60):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for m in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(dxa))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def write_cell(cell, text, *, bold=False, size=9, color=TEXT, align="left", fill=None):
    if fill:
        shade_cell(cell, fill)
    set_cell_borders(cell)
    set_cell_margins(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_meta_line(doc, label, value, url=None, success=False):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=1.15)
    r = p.add_run(label + "  ")
    set_run_font(r, size=11, bold=True, color=TEXT)
    if url:
        add_hyperlink(p, value, url, size=11, bold=True)
    else:
        color = GREEN if success else TEXT
        r2 = p.add_run(value)
        set_run_font(r2, size=11, bold=success, color=color)


def add_body(doc, text, *, before=0, after=8, color=TEXT):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after)
    add_mixed(p, text, color=color)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=4)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    r = p.add_run("•  ")
    set_run_font(r, size=11, color=MUTED)
    add_mixed(p, text)


def add_bullet_link(doc, label, url):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=2, line=1.15)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    r = p.add_run("•  ")
    set_run_font(r, size=11, color=MUTED)
    add_hyperlink(p, label, url, size=11)


def add_glance_table(doc, headers, rows, *, right_align_from=2, passed_col=None):
    """headers: list[str]; rows: list[tuple[str, ...]]. Numeric columns right-aligned from index."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    aligns = ["left"] * len(headers)
    aligns[0] = "center"
    for i in range(right_align_from, len(headers)):
        aligns[i] = "right"
    if passed_col is not None:
        aligns[passed_col] = "center"
    for i, h in enumerate(headers):
        write_cell(table.rows[0].cells[i], h, bold=True, size=8, align=aligns[i], fill=HEADER_FILL)
    remaining_col = passed_col - 1 if passed_col else None
    for r_i, row in enumerate(rows):
        fill = ROW_ALT if r_i % 2 == 1 else "FFFFFF"
        for c_i, val in enumerate(row):
            extra_fill = fill
            extra_color = TEXT
            extra_bold = remaining_col is not None and c_i == remaining_col
            if passed_col is not None and c_i == passed_col:
                extra_bold = True
                extra_color = GREEN
                extra_fill = PASSED_FILL
            write_cell(
                table.rows[r_i + 1].cells[c_i],
                val,
                size=8,
                align=aligns[c_i],
                fill=extra_fill,
                bold=extra_bold,
                color=extra_color,
            )
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)
    return table
