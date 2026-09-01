#!/usr/bin/env python3
"""Split existing 80x25k POD mass-import xlsx files into 200x10k files."""
from __future__ import annotations

import time
from pathlib import Path

import openpyxl
import xlsxwriter

SRC_DIR = Path("/Users/lukachrikishvili/Desktop/mass imports/2 million pod")
ROWS_PER_OUT = 10_000
SRC_COUNT = 80
OUT_COUNT = 200


def write_chunk(out_path: Path, headers: list, rows: list) -> None:
    workbook = xlsxwriter.Workbook(
        str(out_path),
        {"constant_memory": True, "strings_to_urls": False},
    )
    sheet = workbook.add_worksheet("Sheet1")
    sheet.write_row(0, 0, headers)
    for i, row in enumerate(rows, start=1):
        sheet.write_row(i, 0, list(row))
    workbook.close()


def main() -> None:
    started = time.time()
    src_files = [
        SRC_DIR / f"pod-mass-import-dev2-34493-part-{i:02d}-of-{SRC_COUNT:02d}.xlsx"
        for i in range(1, SRC_COUNT + 1)
    ]
    missing = [p.name for p in src_files if not p.exists()]
    if missing:
        raise SystemExit(f"Missing source files: {missing}")

    headers = None
    chunk: list = []
    out_part = 1
    for src in src_files:
        t0 = time.time()
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
        ws = wb["Sheet1"]
        it = ws.iter_rows(values_only=True)
        file_headers = list(next(it))
        if headers is None:
            headers = file_headers
        elif file_headers != headers:
            raise SystemExit(f"Header mismatch in {src.name}")
        for row in it:
            chunk.append(row)
            if len(chunk) == ROWS_PER_OUT:
                out_path = SRC_DIR / f"pod-mass-import-dev2-34493-part-{out_part:03d}-of-{OUT_COUNT:03d}.xlsx"
                write_chunk(out_path, headers, chunk)
                size_mb = out_path.stat().st_size / (1024 * 1024)
                print(
                    f"WROTE {out_path.name} rows={len(chunk)} size_mb={size_mb:.1f} from={src.name}",
                    flush=True,
                )
                out_part += 1
                chunk = []
        wb.close()
        print(f"READ {src.name} elapsed_s={time.time() - t0:.1f}", flush=True)

    if chunk:
        raise SystemExit(f"Leftover rows={len(chunk)}; expected exact 10k chunks")
    if out_part != OUT_COUNT + 1:
        raise SystemExit(f"Expected {OUT_COUNT} outputs, wrote {out_part - 1}")

    for src in src_files:
        src.unlink()
        print(f"REMOVED {src.name}", flush=True)

    print(
        f"DONE files={OUT_COUNT} rows_each={ROWS_PER_OUT} total_rows={OUT_COUNT * ROWS_PER_OUT} elapsed_s={time.time() - started:.1f}"
    )


if __name__ == "__main__":
    main()
