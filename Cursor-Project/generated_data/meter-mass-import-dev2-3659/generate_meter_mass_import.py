#!/usr/bin/env python3
"""Generate Dev2 meter mass-import files cloned from meter id 3659.

One meter per G2M POD from POD parts 001-050. Meter_id empty => create.
Grid, installment/remove dates, and scales match meter 3659.
Meter number is unique (G2MM + POD 7-digit sequence) because the same
number + grid + overlapping dates cannot be installed on many PODs.
"""
from __future__ import annotations

import json
import os
import ssl
import time
import urllib.error
import urllib.request
from datetime import date, datetime
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

ENV_PATH = Path("/Users/lukachrikishvili/Desktop/CURSOR-PROJECT/Cursor-Project/.env")
SCRIPT_DIR = Path(__file__).resolve().parent
POD_DIR = Path("/Users/lukachrikishvili/Desktop/mass imports/2 million pod")
OUT_DIR = Path("/Users/lukachrikishvili/Desktop/mass imports/500k meters")
TEMPLATE_PATH = SCRIPT_DIR / "official-meter-mass-import-template.xlsx"
USER_TEMPLATE = Path("/Users/lukachrikishvili/Downloads/meters-mass-import-template.xlsx")

BASE = "https://devapps.energo-pro.bg/backend/phoenix2-dev"
TEMPLATE_URL = f"{BASE}/mass-import/METERS/template/download"
POD_FILE_COUNT = 50
ROWS_PER_FILE = 10_000
GRID_OPERATOR = "BIG DATA"
SELECT_ALL = "YES"
INSTALLMENT = date(2000, 12, 1)
REMOVE = date(2040, 4, 1)
# Already has clone-source meter 3659 (BIGDATA5453242) on the same dates.
SKIP_POD_IDENTIFIER = "G2M0493944LI0SGN32C8LBKV8KUF4Z86I"
CTX = ssl._create_unverified_context()


def load_env() -> None:
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.strip().startswith("#"):
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def login() -> str:
    body = json.dumps(
        {"user": os.environ["PORTAL_USER"], "password": os.environ["PASSWORD"]}
    ).encode()
    req = urllib.request.Request(
        os.environ["DEVAUTHAPI"],
        data=body,
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(req, timeout=30, context=CTX) as resp:
        payload = json.load(resp)
    token = payload.get("jwt")
    if not token:
        raise SystemExit("LOGIN_FAILED: no jwt")
    return token


def download_live_template(token: str) -> list[str]:
    req = urllib.request.Request(
        TEMPLATE_URL,
        headers={"Authorization": f"Bearer {token}"},
    )
    with urllib.request.urlopen(req, timeout=120, context=CTX) as resp:
        data = resp.read()
    TEMPLATE_PATH.write_bytes(data)
    wb = load_workbook(TEMPLATE_PATH, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    if headers[:7] != [
        "Meter_id",
        "Meter_number",
        "Grid_operator",
        "Meter_POD",
        "Meter_installment_date",
        "Meter_remove_date",
        "Meter_tariff_sellect_all",
    ]:
        raise SystemExit(f"UNEXPECTED_CORE_HEADERS {headers[:7]}")
    return headers


def compare_user_template(live_headers: list[str]) -> None:
    if not USER_TEMPLATE.exists():
        print("USER_TEMPLATE_MISSING", flush=True)
        return
    wb = load_workbook(USER_TEMPLATE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    user_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    print(
        f"TEMPLATE_COMPARE live={len(live_headers)} user={len(user_headers)} equal={user_headers == live_headers}",
        flush=True,
    )


def meter_number_for(pod_identifier: str) -> str:
    seq = pod_identifier[3:10]
    if len(seq) != 7 or not seq.isdigit():
        raise SystemExit(f"BAD_POD_IDENTIFIER {pod_identifier}")
    return f"G2MM{seq}"


def iter_pod_identifiers() -> list[str]:
    identifiers: list[str] = []
    skipped = 0
    for part in range(1, POD_FILE_COUNT + 1):
        path = POD_DIR / f"pod-mass-import-dev2-34493-part-{part:03d}-of-200.xlsx"
        if not path.exists():
            raise SystemExit(f"MISSING {path}")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(min_row=2, values_only=True)
        part_ids: list[str] = []
        for row in rows:
            if not row or row[3] is None:
                continue
            ident = str(row[3]).strip()
            if ident == SKIP_POD_IDENTIFIER:
                skipped += 1
                continue
            part_ids.append(ident)
        wb.close()
        print(f"READ {path.name} pods={len(part_ids)} skipped={skipped}", flush=True)
        identifiers.extend(part_ids)
    if skipped != 1:
        raise SystemExit(f"EXPECTED_SKIP_1 got={skipped}")
    return identifiers


def write_file(part: int, file_count: int, chunk: list[str], headers: list[str]) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"meter-mass-import-dev2-3659-part-{part:03d}-of-{file_count:03d}.xlsx"
    workbook = xlsxwriter.Workbook(
        str(path),
        {"constant_memory": True, "strings_to_urls": False},
    )
    sheet = workbook.add_worksheet("Sheet1")
    date_fmt = workbook.add_format({"num_format": "dd.mm.yyyy"})
    installment_dt = datetime(INSTALLMENT.year, INSTALLMENT.month, INSTALLMENT.day)
    remove_dt = datetime(REMOVE.year, REMOVE.month, REMOVE.day)
    sheet.write_row(0, 0, headers)
    for offset, pod_identifier in enumerate(chunk):
        excel_row = offset + 1
        sheet.write_string(excel_row, 1, meter_number_for(pod_identifier))
        sheet.write_string(excel_row, 2, GRID_OPERATOR)
        sheet.write_string(excel_row, 3, pod_identifier)
        sheet.write_datetime(excel_row, 4, installment_dt, date_fmt)
        sheet.write_datetime(excel_row, 5, remove_dt, date_fmt)
        sheet.write_string(excel_row, 6, SELECT_ALL)
    workbook.close()
    return path


def main() -> None:
    started = time.time()
    load_env()
    token = login()
    print("LOGIN_OK", flush=True)
    headers = download_live_template(token)
    print(f"LIVE_TEMPLATE cols={len(headers)} path={TEMPLATE_PATH}", flush=True)
    compare_user_template(headers)
    identifiers = iter_pod_identifiers()
    print(f"PODS_TO_METER {len(identifiers)}", flush=True)
    file_count = (len(identifiers) + ROWS_PER_FILE - 1) // ROWS_PER_FILE
    for part in range(1, file_count + 1):
        start = (part - 1) * ROWS_PER_FILE
        chunk = identifiers[start : start + ROWS_PER_FILE]
        t0 = time.time()
        path = write_file(part, file_count, chunk, headers)
        size_mb = path.stat().st_size / (1024 * 1024)
        print(
            f"WROTE {path.name} rows={len(chunk)} size_mb={size_mb:.2f} elapsed_s={time.time() - t0:.1f}",
            flush=True,
        )
    print(
        f"DONE files={file_count} total_rows={len(identifiers)} elapsed_s={time.time() - started:.1f}",
        flush=True,
    )


if __name__ == "__main__":
    main()
