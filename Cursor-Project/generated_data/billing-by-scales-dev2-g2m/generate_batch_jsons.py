#!/usr/bin/env python3
"""Generate Dev2 billing-by-scales/batch JSON files (max 7000 requests each; Dev2 heap limit)."""
from __future__ import annotations

import json
import time
from pathlib import Path

from openpyxl import load_workbook

POD_DIR = Path("/Users/lukachrikishvili/Desktop/mass imports/2 million pod")
OUT_DIR = Path(__file__).resolve().parent.parent / "500k billing-by-scales"
POD_FILE_COUNT = 50
ROWS_PER_FILE = 7_000
SKIP = {
    "G2M000000165J9K26TYA0L9PXQ2OXBGDQ",  # already created via POST /billing-by-scales
    "G2M000000204GZ93GQCRC5FYOMC97L4CU",  # already created via /batch
    "G2M0493944LI0SGN32C8LBKV8KUF4Z86I",  # clone-source meter 3659, no G2MM meter
}
PERIODS = [
    ("2026-07-01", "2026-07-06"),
    ("2026-07-07", "2026-07-12"),
    ("2026-07-13", "2026-07-18"),
    ("2026-07-19", "2026-07-24"),
    ("2026-07-25", "2026-07-31"),
]


def meter_number_for(pod_identifier: str) -> str:
    return f"G2MM{pod_identifier[3:10]}"


def build_rows(meter: str) -> list[dict]:
    rows: list[dict] = []
    index = 0
    for i, (start, end) in enumerate(PERIODS):
        rows.append(
            {
                "periodFrom": start,
                "periodTo": end,
                "meterNumber": meter,
                "scaleNumber": "1",
                "scaleCode": "BIG SCALE CODE",
                "scaleType": "BIG DATA TYPE",
                "newMeterReading": (i + 1) * 10,
                "oldMeterReading": i * 10,
                "difference": 10,
                "multiplier": 1,
                "totalVolumes": 10,
                "index": index,
            }
        )
        index += 1
        rows.append(
            {
                "periodFrom": start,
                "periodTo": end,
                "meterNumber": meter,
                "scaleType": "BIG",
                "tariffScale": "BIG DATA TARIFF",
                "volumes": 100,
                "unitPrice": 1,
                "totalValue": 100,
                "index": index,
            }
        )
        index += 1
    return rows


def build_request(pod: str) -> dict:
    meter = meter_number_for(pod)
    return {
        "identifier": pod,
        "dateFrom": "2026-07-01",
        "dateTo": "2026-07-31",
        "billingPowerInKw": 2,
        "invoiceNumber": meter,
        "invoiceDate": "2026-07-01T00:00:00.000Z",
        "invoiceCorrection": None,
        "correction": False,
        "override": False,
        "billingByScalesTableCreateRequests": build_rows(meter),
        "saveRecordForIntermediatePeriod": True,
        "saveRecordForMeterReadings": True,
    }


def iter_pods() -> list[str]:
    identifiers: list[str] = []
    skipped = 0
    for part in range(1, POD_FILE_COUNT + 1):
        path = POD_DIR / f"pod-mass-import-dev2-34493-part-{part:03d}-of-200.xlsx"
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[3] is None:
                continue
            ident = str(row[3]).strip()
            if ident in SKIP:
                skipped += 1
                continue
            identifiers.append(ident)
        wb.close()
        print(f"READ {path.name} total={len(identifiers)} skipped={skipped}", flush=True)
    return identifiers


def main() -> None:
    started = time.time()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    identifiers = iter_pods()
    file_count = (len(identifiers) + ROWS_PER_FILE - 1) // ROWS_PER_FILE
    print(f"PODS {len(identifiers)} files={file_count}", flush=True)
    for part in range(1, file_count + 1):
        start = (part - 1) * ROWS_PER_FILE
        chunk = identifiers[start : start + ROWS_PER_FILE]
        payload = {"requests": [build_request(pod) for pod in chunk]}
        path = OUT_DIR / f"billing-by-scales-batch-dev2-part-{part:03d}-of-{file_count:03d}.json"
        t0 = time.time()
        path.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")
        size_mb = path.stat().st_size / (1024 * 1024)
        print(
            f"WROTE {path.name} requests={len(chunk)} size_mb={size_mb:.1f} elapsed_s={time.time() - t0:.1f}",
            flush=True,
        )
    print(f"DONE files={file_count} requests={len(identifiers)} elapsed_s={time.time() - started:.1f}", flush=True)


if __name__ == "__main__":
    main()
