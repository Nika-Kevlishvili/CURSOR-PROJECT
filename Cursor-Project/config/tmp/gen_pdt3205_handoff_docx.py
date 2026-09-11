# -*- coding: utf-8 -*-
"""Generate PDT-3205 Virtual POS testing handoff Word document."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

out = Path(r"c:\Users\N.kevlishvili\Cursor\Cursor-Project\reports\Chat reports\2026\August\21") / "PDT-3205_Virtual_POS_Testing_Handoff.docx"
out.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    return para

def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")

# Title
title = doc.add_heading("PDT-3205 — Virtual POS Testing Handoff", 0)
p("Purpose: Full chat-session handoff so a new tester can continue Virtual POS / EasyPay copy testing in a new Cursor chat.")
p("Prepared for: Mariam Kvitsiani (for assignment to new tester)")
p("Prepared by: Nika Kevlishvili (QA) — Cursor session handoff")
p("Date: 2026-08-21")
p("Environment used: Dev (Phoenix API http://10.236.20.11:8091, payment-api http://10.236.20.11:9091)")

h("1. Ticket and goal")
bullets([
    "Jira parent: PDT-3205 — Copy EasyPay functionality in another collection channel (Virtual POS)",
    "URL: https://oppa-support.atlassian.net/browse/PDT-3205",
    "Requirement: Copy ALL EasyPay endpoints with the exact same logic; ONLY allowed difference is payments bind to VirtualPos collection channel (not EasyPay).",
    "Status during testing: Testing",
    "GE / Confluence: Online payment easy pay final — page ID 259162123 — https://asterbit.atlassian.net/wiki/spaces/Phoenix/pages/259162123/Online+payment+easy+pay+final",
    "Also: Phase 2 Online payment Virtual POS terminal — 611024897",
])

h("2. How to continue in a new Cursor chat")
bullets([
    "Open a new chat and attach this Word file (or paste the path below).",
    "Say: Continue PDT-3205 Virtual POS testing on Dev from this handoff. Do not re-file PDT-3369/3370/3371/3372 unless re-verified after a fix.",
    "Workspace: Cursor-Project; EnergoTS branch: cursor; Phoenix aligned to origin/dev for payment-api.",
    "Playwright suite: Cursor-Project/EnergoTS/tests/cursor/PDT-3205-copy-easypay-to-virtual-pos.spec.ts (+ pdt-3205-*.fixtures.ts)",
    "Backend test cases MD: Cursor-Project/test_cases/Backend/Copy_EasyPay_to_Virtual_POS.md",
    "Always run Playwright with: BASE_URL=http://10.236.20.11:8091 (classic Dev). Default phoenix2 BASE_URL hits :9092 where VPOS merchant may be 7000005 and causes systemic STATUS 96.",
])

h("3. Protocol facts (payment-api)")
bullets([
    "Virtual POS paths: GET /virtualpos/init-pay, /confirm-pay, /calculate-check-sum-init, /calculate-check-sum-confirm",
    "EasyPay paths: GET /epay/... (same shape)",
    "NOT in Phoenix core Swagger — use payment-api Swagger: http://10.236.20.11:9091/swagger-ui/index.html",
    "Dev merchants: EasyPay 7000005; Virtual POS 7000006 (same HMAC secret currently on Dev — isolation gap)",
    "Channel names: EasyPay / VirtualPos (combineLiabilities true/false matters for INVOICES)",
    "Amounts in stotinki (coins). DATE/TID format yyyyMMddHHmmss (+ STAN 077118 + AID 100020 for TID from test helpers).",
    "HTTP 200 with protocol STATUS in JSON body: 00 OK, 13 invalid amount, 14 invalid subscriber, 62 no obligation, 80 temp unable, 93 bad checksum, 94 repeat, 96 general error.",
    "These are GET + query params — there is NO JSON request body. Swagger Try it out = query fields.",
])

h("4. What was already delivered")
h("4.1 Automation / docs", 2)
bullets([
    "Backend TCs generated (Rule 35 pipeline earlier in session).",
    "Playwright PDT-3205 suite (~14 cases) — last full run ~11 passed / 3 failed (product gaps).",
    "Senior QA findings vs Confluence + EasyPay code comparison.",
])

h("4.2 Internal Bugs filed under PDT-3205", 2)
bullets([
    "PDT-3369 — Confirm DATE ignored for accounting period / payment package (LocalDate.now + current AP). Highest. https://oppa-support.atlassian.net/browse/PDT-3369",
    "PDT-3370 — Proforma/null dueDate → per-invoice VALIDTO \"0\" (EasyPay uses today+30). Highest. https://oppa-support.atlassian.net/browse/PDT-3370",
    "PDT-3371 — LONGDESC missing invoice document number/date (no findInvoiceDates enrichment). High. https://oppa-support.atlassian.net/browse/PDT-3371",
    "PDT-3372 — Partial INVOICES accepts liability IDs missing from init snapshot (no EasyPay size==size check). Highest. https://oppa-support.atlassian.net/browse/PDT-3372",
])
p("Assignee on bugs: Giorgi Guliashvili. Tester/Reporter: Nika Kevlishvili. Environment: Dev.")

h("5. Playwright results snapshot (Dev, BASE_URL 8091)")
bullets([
    "PASSED examples: TC-BE-1 CHECK, TC-BE-2 happy path channel binding, TC-BE-8 immediate offset after confirm (runtime passed despite async code), TC-BE-10..12, 14, 17..20 (14/62/93/94/96 negatives).",
    "FAILED (match filed bugs): TC-BE-5 (DATE/AP → HTTP 400), TC-BE-6 (VALIDTO \"0\"), TC-BE-7 (LONGDESC empty Фактура).",
    "Command: cd Cursor-Project/EnergoTS; $env:BASE_URL='http://10.236.20.11:8091'; npx playwright test tests/cursor/PDT-3205-copy-easypay-to-virtual-pos.spec.ts --workers=1",
])

h("6. Bug details for retest after fix")

h("6.1 PDT-3369 — DATE / AP / package", 2)
bullets([
    "Confirm with historical DATE in prior OPEN AP (e.g. DATE=20260615120000) while current AP is August (1044).",
    "VPOS uses current AP + package date LocalDate.now() → APPLICATION_ERROR 400 mixing June paymentDate with August package.",
    "EasyPay uses paymentDate + fetchAccountPeriodIdForDateWithFallback — does NOT do this.",
    "Example confirm that failed: TID 20260821171137077118100020, IDN 6000150093, TOTAL 5113, DATE 20260615120000 → 400.",
])

h("6.2 PDT-3370 — VALIDTO \"0\"", 2)
bullets([
    "combineLiabilities=false; proforma liability with null dueDate; BILLING init.",
    "Root VALIDTO may be today+30; INVOICES[].VALIDTO = \"0\".",
    "Example customer 6000150094 / liability 89501.",
])

h("6.3 PDT-3371 — LONGDESC empty invoice part", 2)
bullets([
    "After issue-invoice, init should show Фактура: {doc}/{date}.",
    "Actual: Фактура: ; — VPOS fetchLiabilities skips invoice-date enrichment present in EasyPayBaseService.",
    "Example: IDN 6000150099, expected Фактура: 1000069219/21.08.2026.",
])

h("6.4 PDT-3372 — Partial INVOICES size check", 2)
bullets([
    "combine=false; init with liabilities A+B; create C after init; confirm TOTAL=A coins, INVOICES=A,C.",
    "VPOS STATUS 00 and pays only A; EasyPay code rejects size mismatch.",
    "Evidence: customer 6084552 / 6000150106; A=89514 B=89515 C=89516; payment 27097; TID 20260821182252077118100020; TOTAL 2556.",
    "Portals: customer https://devapps.energo-pro.bg/app/phoenix1-dev/customers/preview/basic?id=6084552 payment https://devapps.energo-pro.bg/app/phoenix1-dev/payments/preview?id=27097",
    "Note: EasyPay live side-by-side not run; parity expected from code.",
])

h("7. Candidate gaps NOT filed yet (need more verification)")
bullets([
    "#1 Async offsetting (CompletableFuture.runAsync) — TC-BE-8 PASSED on Dev (liability already 0 right after confirm). Code still async; do not file unless race reproduced.",
    "#3 LPF fallback / broader LPF trigger vs EasyPay — NEEDS_MORE_SETUP (overdue + LPF data).",
    "#4 Cover date param vs EasyPay (null vs initDate) — incomplete without EasyPay A/B DB compare.",
    "Shared EasyPay+VPOS issues (IDN drops billing group on init response; re-init same TID stale snapshot; principal-increased covering vs Confluence; public checksum endpoints) — product decision whether to file under PDT-3205 or EasyPay tickets.",
])

h("8. External stakeholder — Valeri (#production-fixes)")
bullets([
    "Thread: https://asterbit.slack.com/archives/C0A2YDXTJKU/p1786950996.375319 (PDT-3205)",
    "Valeri got STATUS 96 on Dev using MERCHANTID=7000005 on /virtualpos.",
    "Root cause: wrong merchant. Same IDN/TID with 7000006 → STATUS 00.",
    "Nika drafted Slack reply explaining 7000006 for Virtual POS vs 7000005 for EasyPay.",
])

h("9. Manual happy-path reminder")
bullets([
    "1) calculate-check-sum-init TYPE=CHECK (no TID) → 2) init-pay CHECK → 3) checksum BILLING+TID → 4) init BILLING → 5) checksum confirm with DATE+TOTAL → 6) confirm-pay.",
    "Never change DATE/TOTAL/TID between checksum-confirm and confirm-pay (else STATUS 93).",
    "DATE must not be in the future (else 96).",
])

h("10. Suggested next steps for new tester")
bullets([
    "1. Re-run full PDT-3205 Playwright on Dev with BASE_URL=8091; confirm same 3 product failures until fixes land.",
    "2. After each bug fix (3369–3372), retest that ticket with listed evidence IDs/payloads.",
    "3. Optionally run EasyPay twin of PDT-3372 scenario to confirm EasyPay rejects.",
    "4. Finish verification of LPF (#3) and cover-date (#4) before filing more bugs.",
    "5. Keep helping Valeri with correct merchant/config for UAT/prod planning.",
    "6. Do not edit Phoenix code from Cursor (Tier A). EnergoTS tests only via energo-ts-test workflow.",
])

h("11. Key file paths")
bullets([
    "Spec: Cursor-Project/EnergoTS/tests/cursor/PDT-3205-copy-easypay-to-virtual-pos.spec.ts",
    "Fixtures: Cursor-Project/EnergoTS/tests/cursor/pdt-3205-copy-easypay-to-virtual-pos.fixtures.ts",
    "TCs: Cursor-Project/test_cases/Backend/Copy_EasyPay_to_Virtual_POS.md",
    "Bug reviews: Cursor-Project/reports/Bug Reports/2026/august/21/BugReview_*.md",
    "This handoff: " + str(out),
])

h("12. Agents / rules used in original session")
bullets([
    "PhoenixExpert, Senior QA, EnvironmentResolver, Test case / Playwright / Bug reporter workflows",
    "DB: Dev PostgreSQL for evidence (virtual_pos_payments, customer_payments, etc.)",
    "Confluence/Jira: REST fallback when MCP unavailable",
])

p("")
p("End of handoff — continue testing from section 10.", bold=True)

doc.save(str(out))
print(out)
print("bytes", out.stat().st_size)
