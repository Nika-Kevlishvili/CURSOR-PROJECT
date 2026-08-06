"""Build separate CREATE and EDIT POD mass-import Excel files for Dev (PDT-3143)."""
import json
import random
import string
import urllib.request
import uuid
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(r"c:\Users\N.kevlishvili\Cursor\Cursor-Project\config\templates\pod-mass-import")
TOKEN = json.loads(
    Path(r"c:\Users\N.kevlishvili\Cursor\Cursor-Project\EnergoTS\fixtures\token.json").read_text(
        encoding="utf-8"
    )
)["token"]
API = "http://10.236.20.11:8091"


def download_template() -> Path:
    out = BASE / "pods_official_template_dev.xlsx"
    req = urllib.request.Request(f"{API}/mass-import/PODS/template/download", method="GET")
    req.add_header("Authorization", f"Bearer {TOKEN}")
    with urllib.request.urlopen(req, timeout=60) as resp:
        out.write_bytes(resp.read())
    return out


def set_str(ws, row: int, col0: int, value):
    cell = ws.cell(row=row, column=col0 + 1)
    cell.value = str(value)


def fill_create_row(ws, create_id: str, energy_sharing: str | None):
    """Excel row 2 — CREATE (POD_id blank)."""
    set_str(ws, 2, 3, create_id)
    set_str(ws, 2, 5, "PDT3143 CREATE POD")
    set_str(ws, 2, 6, "ЕСО ЕАД")
    ws.cell(2, 9).value = 1
    set_str(ws, 2, 12, "CONSUMER")
    set_str(ws, 2, 13, "NON_HOUSEHOLD")
    set_str(ws, 2, 14, "LOW")
    set_str(ws, 2, 15, "SETTLEMENT_PERIOD")
    set_str(ws, 2, 19, "NO")
    set_str(ws, 2, 20, "standard country")
    set_str(ws, 2, 21, "standard region")
    set_str(ws, 2, 22, "standard municipality")
    set_str(ws, 2, 23, "standard populated place")
    set_str(ws, 2, 24, "standard zip code")
    if energy_sharing:
        set_str(ws, 2, 52, energy_sharing)  # Energy sharing model


def fill_edit_row(ws, identifier: str, energy_sharing: str | None, version_action: str = "E"):
    """Excel row 2 — EDIT only file (POD_id non-blank)."""
    set_str(ws, 2, 0, "1")
    set_str(ws, 2, 2, version_action)
    set_str(ws, 2, 3, identifier)
    set_str(ws, 2, 15, "SETTLEMENT_PERIOD")
    if energy_sharing:
        set_str(ws, 2, 52, energy_sharing)


def clear_data_rows(ws):
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


def build_file(template: Path, name: str, filler) -> Path:
    wb = load_workbook(template)
    ws = wb.active
    assert ws.title == "Sheet1", ws.title
    clear_data_rows(ws)
    filler(ws)
    # ensure header count still 53
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert headers[52] is not None, "col 52 header missing"
    out = BASE / name
    wb.save(out)
    # verify
    wb2 = load_workbook(out)
    ws2 = wb2.active
    nonempty = {c: ws2.cell(2, c).value for c in range(1, 54) if ws2.cell(2, c).value is not None}
    meta = {"file": str(out), "row2_nonempty_cols": nonempty, "header_52": ws2.cell(1, 53).value}
    (BASE / (name + ".meta.json")).write_text(
        json.dumps(meta, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    return out


def upload(path: Path) -> int:
    boundary = "----CursorBoundary" + uuid.uuid4().hex
    data = path.read_bytes()
    parts = [
        f"--{boundary}".encode(),
        f'Content-Disposition: form-data; name="file"; filename="{path.name}"'.encode(),
        b"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        b"",
        data,
        f"--{boundary}--".encode(),
        b"",
    ]
    payload = b"\r\n".join(parts)
    req = urllib.request.Request(
        f"{API}/mass-import/PODS/files/upload", data=payload, method="POST"
    )
    req.add_header("Authorization", f"Bearer {TOKEN}")
    req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            return resp.status
    except urllib.error.HTTPError as e:
        print("UPLOAD ERR", e.code, e.read()[:800])
        return e.code


def main():
    BASE.mkdir(parents=True, exist_ok=True)
    tpl = download_template()
    wb = load_workbook(tpl)
    ws = wb.active
    print("template header52=", ws.cell(1, 53).value)

    suffix = "".join(random.choices(string.ascii_uppercase + string.digits, k=18))
    create_id = ("32XMI" + suffix)[:33]
    (BASE / "last-create-identifier.txt").write_text(create_id, encoding="utf-8")

    create_path = build_file(
        tpl,
        "pods_mass_import_CREATE_only_pdt3143.xlsx",
        lambda w: fill_create_row(w, create_id, "ENERGY_GROUP"),
    )
    edit_path = build_file(
        tpl,
        "pods_mass_import_EDIT_only_pdt3143.xlsx",
        lambda w: fill_edit_row(w, "32XOVMPIIUGMY7318258694", "ENERGY_COMMUNITY", "E"),
    )

    # remove old combined file if present
    combined = BASE / "pods_mass_import_minimal_create_edit.xlsx"
    if combined.exists():
        combined.unlink()
        print("removed combined file")

    print("CREATE file", create_path, "id", create_id)
    print("EDIT file", edit_path)
    print("UPLOAD CREATE", upload(create_path))
    print("UPLOAD EDIT", upload(edit_path))


if __name__ == "__main__":
    main()
