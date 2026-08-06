import json
import urllib.request
import uuid
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

BASE = Path(r"c:\Users\N.kevlishvili\Cursor\Cursor-Project\config\templates\pod-mass-import")
TOKEN = json.loads(
    Path(r"c:\Users\N.kevlishvili\Cursor\Cursor-Project\EnergoTS\fixtures\token.json").read_text(
        encoding="utf-8"
    )
)["token"]
API = "http://10.236.20.11:8091"

wb = load_workbook(BASE / "pods_official_template_dev.xlsx", read_only=True)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, 54)]
wb.close()
print("HEADER_52=", headers[52])

create_id = (BASE / "last-create-identifier.txt").read_text(encoding="utf-8").strip()
edit_row = {
    0: "1",
    2: "E",
    3: create_id,
    5: "PDT3143 EDIT POD RERUN",
    15: "SETTLEMENT_PERIOD",
    19: "NO",
    20: "standard country",
    21: "standard region",
    22: "standard municipality",
    23: "standard populated place",
    24: "standard zip code",
    52: "ENERGY_COMMUNITY",
}
path = BASE / "pods_mass_import_EDIT_only_pdt3143.xlsx"
wb2 = xlsxwriter.Workbook(str(path), {"strings_to_urls": False})
ws2 = wb2.add_worksheet("Sheet1")
for i, h in enumerate(headers):
    ws2.write_string(0, i, str(h) if h is not None else "")
for col0, val in edit_row.items():
    if isinstance(val, (int, float)):
        ws2.write_number(1, col0, val)
    else:
        ws2.write_string(1, col0, str(val))
wb2.close()

boundary = "----CursorBoundary" + uuid.uuid4().hex
data = path.read_bytes()
parts = [
    f"--{boundary}".encode(),
    f'Content-Disposition: form-data; name="file"; filename="{path.name}"'.encode(),
    b"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    b"",
    data,
    f"--{boundary}--".encode(),
    b"",
]
payload = b"\r\n".join(parts)
req = urllib.request.Request(f"{API}/mass-import/PODS/files/upload", data=payload, method="POST")
req.add_header("Authorization", f"Bearer {TOKEN}")
req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
with urllib.request.urlopen(req, timeout=120) as resp:
    print("UPLOAD EDIT", resp.status, "target", create_id)
