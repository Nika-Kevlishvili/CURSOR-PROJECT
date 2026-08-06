"""Build CREATE/EDIT POD mass-import files with shared strings (POI-compatible)."""
import json
import random
import string
import urllib.request
import uuid
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

BASE = Path(r"c:\Users\N.kevlishvili\Cursor\Cursor-Project\config\templates\pod-mass-import")
TOKEN = json.loads(
    Path(r"c:\Users\N.kevlishvili\Cursor\Cursor-Project\EnergoTS\fixtures\token.json").read_text(
        encoding="utf-8"
    )
)["token"]
API = "http://10.236.20.11:8091"


def download_template_headers() -> list:
    out = BASE / "pods_official_template_dev.xlsx"
    req = urllib.request.Request(f"{API}/mass-import/PODS/template/download", method="GET")
    req.add_header("Authorization", f"Bearer {TOKEN}")
    with urllib.request.urlopen(req, timeout=60) as resp:
        out.write_bytes(resp.read())
    wb = load_workbook(out, read_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, (ws.max_column or 53) + 1)]
    wb.close()
    # trim trailing Nones
    while headers and headers[-1] is None:
        headers.pop()
    return headers


def write_xlsx(path: Path, headers: list, data_row: dict):
    """data_row: 0-based column index -> value"""
    wb = xlsxwriter.Workbook(str(path), {"strings_to_urls": False, "constant_memory": False})
    ws = wb.add_worksheet("Sheet1")
    for i, h in enumerate(headers):
        ws.write_string(0, i, str(h) if h is not None else "")
    for col0, val in data_row.items():
        if val is None:
            continue
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            ws.write_number(1, col0, val)
        else:
            ws.write_string(1, col0, str(val))
    wb.close()


def upload(path: Path) -> int:
    boundary = "----CursorBoundary" + uuid.uuid4().hex
    data = path.read_bytes()
    parts = [
        f"--{boundary}".encode(),
        f'Content-Disposition: form-data; name="file"; filename="{path.name}"'.encode(),
        b"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        b"",
        data,
        f"--{boundary}--".encode(),
        b"",
    ]
    payload = b"\r\n".join(parts)
    req = urllib.request.Request(
        f"{API}/mass-import/PODS/files/upload", data=payload, method="POST"
    )
    req.add_header("Authorization", f"Bearer {TOKEN}")
    req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            print("UPLOAD", path.name, resp.status)
            return resp.status
    except urllib.error.HTTPError as e:
        print("UPLOAD ERR", path.name, e.code, e.read()[:500])
        return e.code


def main():
    BASE.mkdir(parents=True, exist_ok=True)
    headers = download_template_headers()
    print("n_headers", len(headers), "col52", headers[52] if len(headers) > 52 else None)

    suffix = "".join(random.choices(string.ascii_uppercase + string.digits, k=18))
    create_id = ("32XMI" + suffix)[:33]
    (BASE / "last-create-identifier.txt").write_text(create_id, encoding="utf-8")

    create_row = {
        3: create_id,
        5: "PDT3143 CREATE POD",
        6: "ЕСО ЕАД",
        8: 1,
        12: "CONSUMER",
        13: "NON_HOUSEHOLD",
        14: "LOW",
        15: "SETTLEMENT_PERIOD",
        19: "NO",
        20: "standard country",
        21: "standard region",
        22: "standard municipality",
        23: "standard populated place",
        24: "standard zip code",
        52: "ENERGY_GROUP",
    }
    edit_row = {
        0: "1",
        2: "E",
        3: "32XOVMPIIUGMY7318258694",
        15: "SETTLEMENT_PERIOD",
        52: "ENERGY_COMMUNITY",
    }

    create_path = BASE / "pods_mass_import_CREATE_only_pdt3143.xlsx"
    edit_path = BASE / "pods_mass_import_EDIT_only_pdt3143.xlsx"
    write_xlsx(create_path, headers, create_row)
    write_xlsx(edit_path, headers, edit_row)

    print("create_id", create_id)
    upload(create_path)
    upload(edit_path)


if __name__ == "__main__":
    main()
