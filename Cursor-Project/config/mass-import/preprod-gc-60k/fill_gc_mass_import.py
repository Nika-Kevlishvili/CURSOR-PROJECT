"""Fill PreProd Government Compensation mass-import workbook from invoice POD/customer pairs."""
from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
PAIRS_CSV = HERE / "billing-1119-BILLING202601160016-customer-pod-pairs.csv"
TEMPLATE = HERE / "GovernmentCompensation-preprod-template.xlsx"
OUT_XLSX = HERE / "GovernmentCompensation-preprod-60k.xlsx"
OUT_SMOKE = HERE / "GovernmentCompensation-preprod-1k-smoke.xlsx"
RECIPIENT = "2099070601"
CURRENCY = "Евро"
REASON = "PreProd GC mass import 60k"
DOC_DATE = date(2026, 9, 7)
DOC_PERIOD = date(2026, 9, 1)
TARGET_ROWS = 60_000
SMOKE_ROWS = 1_000


def date_cell(ws: Worksheet, value: date) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.number_format = "YYYY-MM-DD"
    return cell


def load_pairs() -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    with PAIRS_CSV.open(encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            customer = (row["customer_identifier"] or "").strip()
            pod = (row["pod_identifier"] or "").strip()
            if not customer or not pod:
                continue
            if customer == RECIPIENT:
                continue
            pairs.append((customer, pod))
    if not pairs:
        raise SystemExit("No customer/POD pairs loaded")
    return pairs


def load_headers() -> tuple[str, list[str]]:
    wb = load_workbook(TEMPLATE, read_only=True)
    sheet = wb.sheetnames[0]
    ws = wb[sheet]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(v) if v is not None else "" for v in header_row]
    wb.close()
    if len(headers) < 11:
        raise SystemExit(f"Unexpected PreProd template columns: {headers}")
    return sheet, headers


def write_workbook(
    path: Path,
    sheet_name: str,
    headers: list[str],
    pairs: list[tuple[str, str]],
    row_count: int,
    number_prefix: str,
    reason: str,
) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    for i in range(row_count):
        customer, pod = pairs[i % len(pairs)]
        ws.append(
            [
                f"{number_prefix}-{i + 1:06d}",
                date_cell(ws, DOC_DATE),
                date_cell(ws, DOC_PERIOD),
                1,
                1,
                reason,
                1,
                CURRENCY,
                customer,
                pod,
                RECIPIENT,
            ]
        )
    wb.save(path)


def main() -> None:
    pairs = load_pairs()
    sheet_name, headers = load_headers()
    write_workbook(OUT_XLSX, sheet_name, headers, pairs, TARGET_ROWS, "GC60K", REASON)
    write_workbook(
        OUT_SMOKE,
        sheet_name,
        headers,
        pairs,
        SMOKE_ROWS,
        "GCSMOKE",
        "PreProd GC mass import smoke 1k",
    )
    print(
        f"pairs={len(pairs)} headers={len(headers)} "
        f"rows={TARGET_ROWS} smoke={SMOKE_ROWS} recipient={RECIPIENT}"
    )
    print(f"out={OUT_XLSX.name}")
    print(f"smoke={OUT_SMOKE.name}")


if __name__ == "__main__":
    main()
