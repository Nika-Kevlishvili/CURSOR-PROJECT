"""Rebuild 5000-row invoice cancellation parts: number without prefix/dash."""
from __future__ import annotations

import math
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import numbers
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "InvoiceCancellationTemplate-preprod.xlsx"
OUT_DIR = HERE / "invoice-cancellation-parts"
CHUNK = 5_000
CANCELLED_NUMBERS = {
    "1100006325",
    "1100014759",
    "1100025620",
    "1100029413",
    "1100031567",
    "1100035540",
}
PREFIX_DASH = re.compile(r"^.*-")


def date_cell(ws: Worksheet, value: date) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.number_format = "YYYY-MM-DD"
    return cell


def number_cell(ws: Worksheet, value: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.number_format = numbers.FORMAT_TEXT
    return cell


def to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        return datetime.strptime(value.strip()[:10], "%Y-%m-%d").date()
    return None


def strip_prefix(raw) -> str:
    text = str(raw).strip() if raw is not None else ""
    if not text:
        return ""
    if "-" in text:
        return PREFIX_DASH.sub("", text)
    return text


def main() -> None:
    wb_t = load_workbook(TEMPLATE, read_only=True)
    sheet_name = wb_t.sheetnames[0]
    headers = list(next(wb_t[sheet_name].iter_rows(min_row=1, max_row=1, values_only=True)))
    wb_t.close()

    sources = sorted(OUT_DIR.glob("InvoiceCancellation-preprod-billing-1119-part-*.xlsx"))
    if not sources:
        raise SystemExit("No existing part files to rebuild")

    rows: list[tuple[str, date]] = []
    seen: set[str] = set()
    skipped: list[str] = []
    for path in sources:
        wb = load_workbook(path, read_only=True)
        for row in wb.active.iter_rows(min_row=2, values_only=True):
            number = strip_prefix(row[0])
            inv_date = to_date(row[1])
            if not number or inv_date is None:
                continue
            if number in CANCELLED_NUMBERS:
                skipped.append(number)
                continue
            if number in seen:
                continue
            seen.add(number)
            rows.append((number, inv_date))
        wb.close()

    for old in sources:
        old.unlink()

    total_parts = max(1, math.ceil(len(rows) / CHUNK))
    for part in range(total_parts):
        chunk = rows[part * CHUNK : (part + 1) * CHUNK]
        out = OUT_DIR / f"InvoiceCancellation-preprod-billing-1119-part-{part + 1:02d}-of-{total_parts:02d}.xlsx"
        wb_out = Workbook(write_only=True)
        ws = wb_out.create_sheet(sheet_name)
        ws.append(headers)
        for number, inv_date in chunk:
            ws.append([number_cell(ws, number), date_cell(ws, inv_date)])
        wb_out.save(out)
        print(f"wrote {out.name} rows={len(chunk)} sample={chunk[0][0]}")

    print(f"kept={len(rows)} skipped_cancelled={len(skipped)} parts={total_parts}")


if __name__ == "__main__":
    main()
