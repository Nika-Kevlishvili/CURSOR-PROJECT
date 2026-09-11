"""Fill PreProd Invoice Cancellation import file from billing 1119 REAL invoices."""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "invoice-cancellation-template-invoice-cancellation_download-template.xlsx"
OUT_XLSX = HERE / "InvoiceCancellation-preprod-billing-1119.xlsx"
OUT_TEMPLATE = HERE / "InvoiceCancellationTemplate-preprod.xlsx"
JSON_PARTS = [
    Path(r"C:\Users\N.kevlishvili\.cursor\projects\c-Users-N-kevlishvili-Cursor\agent-tools\b1d1e6b5-ac4d-4d74-b367-97dc119c31b5.txt"),
    Path(r"C:\Users\N.kevlishvili\.cursor\projects\c-Users-N-kevlishvili-Cursor\agent-tools\4d56e739-45be-424d-9ec8-732bfe80a053.txt"),
    Path(r"C:\Users\N.kevlishvili\.cursor\projects\c-Users-N-kevlishvili-Cursor\agent-tools\88dc2df5-32db-43ee-83e2-70dee852b2a7.txt"),
    Path(r"C:\Users\N.kevlishvili\.cursor\projects\c-Users-N-kevlishvili-Cursor\agent-tools\0f022f15-edd2-48ae-b8f7-f9fc11ca0dbc.txt"),
    Path(r"C:\Users\N.kevlishvili\.cursor\projects\c-Users-N-kevlishvili-Cursor\agent-tools\6eb3b0ca-e77e-43f5-985c-930040d52f20.txt"),
    Path(r"C:\Users\N.kevlishvili\.cursor\projects\c-Users-N-kevlishvili-Cursor\agent-tools\1758fa93-f9dd-4ca0-80c8-7e65e53d49e4.txt"),
    Path(r"C:\Users\N.kevlishvili\.cursor\projects\c-Users-N-kevlishvili-Cursor\agent-tools\a8890bfb-68a3-4abc-be18-a0a30b5dc15b.txt"),
]


def date_cell(ws: Worksheet, value: date) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.number_format = "YYYY-MM-DD"
    return cell


def parse_date(raw: str) -> date:
    return datetime.strptime(raw[:10], "%Y-%m-%d").date()


def main() -> None:
    rows: list[tuple[str, date]] = []
    seen: set[str] = set()
    for path in JSON_PARTS:
        if not path.exists():
            raise SystemExit(f"Missing batch file: {path}")
        data = json.loads(path.read_text(encoding="utf-8"))
        for item in data:
            number = (item.get("invoice_number") or "").strip()
            raw_date = (item.get("invoice_date") or "").strip()
            if not number or not raw_date or number in seen:
                continue
            seen.add(number)
            rows.append((number, parse_date(raw_date)))

    wb_t = load_workbook(TEMPLATE, read_only=True)
    sheet_name = wb_t.sheetnames[0]
    headers = list(next(wb_t[sheet_name].iter_rows(min_row=1, max_row=1, values_only=True)))
    wb_t.close()
    OUT_TEMPLATE.write_bytes(TEMPLATE.read_bytes())

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    for number, inv_date in rows:
        ws.append([number, date_cell(ws, inv_date)])
    wb.save(OUT_XLSX)
    print(f"rows={len(rows)} headers={headers} out={OUT_XLSX.name}")


if __name__ == "__main__":
    main()
